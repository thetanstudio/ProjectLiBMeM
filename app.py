from flask import Flask, render_template, request, redirect, session, url_for 
from flask import session, send_file 
from openpyxl import Workbook
from openpyxl import load_workbook
from flask import send_file
from io import BytesIO
from flask import url_for
import os 
from supabase import create_client, Client
from dotenv import load_dotenv





from datetime import datetime, timedelta , date

os.makedirs("static/uploads", exist_ok=True)
from werkzeug.utils import secure_filename

def has_borrowed(student_name):
    return any(i for i in issues if i["student"] == student_name and not i["returned"])
app = Flask(__name__)
app.secret_key = "libmem_secret"

load_dotenv()


# =========================
# SUPABASE DATABASE
# =========================
SUPABASE_URL = os.environ.get("SUPABASE_URL")
SUPABASE_KEY = os.environ.get("SUPABASE_SECRET_KEY")

supabase: Client = create_client(
    SUPABASE_URL,
    SUPABASE_KEY
)

# =========================
# SUPABASE DATABASE HELPERS
# =========================

def load_books_from_db():
    global books, books_from_excel, books_excel_name

    result = supabase.table("books").select("*").execute()
    rows = result.data or []

    books.clear()

    # Restore Excel mode from Supabase after page changes/restarts
    books_from_excel = (
        bool(rows)
        and all(row.get("source") == "excel" for row in rows)
    )

    books_excel_name = (
        "Excel data saved in Supabase"
        if books_from_excel else None
    )

    for row in rows:
        books.append({
            "id": row["id"],
            "name": row["name"],
            "author": row["author"],
            "publisher": row.get("publisher") or "",
            "isbn": row.get("isbn") or "",
            "price": row.get("price") or "",
            "subtitle": row.get("subtitle") or "",
            "year_published": row.get("year_published") or "",
            "cover": row.get("cover") or ""
        })



def save_book_to_db(book):
    supabase.table("books").insert({
        "id": book["id"],
        "name": book["name"],
        "author": book["author"],
        "publisher": book.get("publisher", ""),
        "isbn": book.get("isbn", ""),
        "price": book.get("price", ""),
        "subtitle": book.get("subtitle", ""),
        "year_published": book.get("year_published", ""),
        "cover": book.get("cover", ""),
        "source": "manual"
    }).execute()


def delete_book_from_db(book_id):
    supabase.table("books").delete().eq(
        "id", book_id
    ).execute()


def update_book_in_db(book):
    supabase.table("books").update({
        "name": book["name"],
        "author": book["author"],
        "publisher": book.get("publisher", ""),
        "isbn": book.get("isbn", ""),
        "price": book.get("price", ""),
        "subtitle": book.get("subtitle", ""),
        "year_published": book.get("year_published", ""),
        "cover": book.get("cover", "")
    }).eq(
        "id", book["id"]
    ).execute()


def get_next_book_id():

    result = (
        supabase
        .table("books")
        .select("id")
        .execute()
    )

    max_number = 0

    for row in result.data:

        book_id = str(row.get("id", ""))

        if book_id.startswith("S") and book_id[1:].isdigit():

            number = int(book_id[1:])

            if number > max_number:
                max_number = number

    return f"S{max_number + 1:04d}"




    # =========================
# STUDENT / CLASS HELPERS
# =========================

def load_students_from_db():
    global students, classes

    student_result = supabase.table("students").select("*").execute()
    class_result = supabase.table("classes").select("*").execute()

    students.clear()
    classes.clear()

    for row in student_result.data:
        students.append({
            "name": row["name"],
            "class": row["class_name"]
        })

    for row in class_result.data:
        classes.append(row["name"])


def save_class_to_db(class_name):
    existing = (
        supabase.table("classes")
        .select("id")
        .eq("name", class_name)
        .execute()
    )

    if not existing.data:
        supabase.table("classes").insert({
            "name": class_name,
            "source": "manual"
        }).execute()


def save_student_to_db(student):
    supabase.table("students").insert({
        "name": student["name"],
        "class_name": student["class"],
        "source": "manual"
    }).execute()


def delete_student_from_db(name, class_name):
    supabase.table("students").delete().eq(
        "name", name
    ).eq(
        "class_name", class_name
    ).execute()


def delete_class_from_db(class_name):
    supabase.table("students").delete().eq(
        "class_name", class_name
    ).execute()

    supabase.table("classes").delete().eq(
        "name", class_name
    ).execute()


def update_student_in_db(old_name, class_name, new_name):
    supabase.table("students").update({
        "name": new_name
    }).eq(
        "name", old_name
    ).eq(
        "class_name", class_name
    ).execute()


def update_class_in_db(old_name, new_name):
    supabase.table("students").update({
        "class_name": new_name
    }).eq(
        "class_name", old_name
    ).execute()

    supabase.table("classes").update({
        "name": new_name
    }).eq(
        "name", old_name
    ).execute()


def delete_history_from_db(option):

    from datetime import datetime, timedelta

    # ==========================================
    # DELETE ALL COMPLETED HISTORY
    # ==========================================

    if option == "all":

        # Delete returned loans ONLY
        supabase.table("loans") \
            .delete() \
            .eq("returned", True) \
            .execute()

        # Delete expired reservations ONLY
        supabase.table("reservations") \
            .delete() \
            .eq("expired", True) \
            .execute()

        return


    # ==========================================
    # DELETE HISTORY OLDER THAN X DAYS
    # ==========================================

    days_map = {
        "5": 5,
        "10": 10,
        "20": 20,
        "30": 30
    }

    days = days_map.get(option)

    if not days:
        return


    cutoff = (
        datetime.now() -
        timedelta(days=days)
    ).strftime("%Y-%m-%d")


    # ==========================================
    # RETURNED LOANS
    # ==========================================

    loans_result = (
        supabase
        .table("loans")
        .select("*")
        .eq("returned", True)
        .execute()
    )

    for loan in loans_result.data:

        return_date = loan.get("return_date")

        if return_date and return_date <= cutoff:

            supabase.table("loans") \
                .delete() \
                .eq("id", loan["id"]) \
                .execute()


    # ==========================================
    # EXPIRED RESERVATIONS
    # ==========================================

    reservations_result = (
        supabase
        .table("reservations")
        .select("*")
        .eq("expired", True)
        .execute()
    )

    for reservation in reservations_result.data:

        expiry_date = reservation.get("reserved_until")

        if expiry_date and expiry_date <= cutoff:

            supabase.table("reservations") \
                .delete() \
                .eq("id", reservation["id"]) \
                .execute()


def load_history_from_db():
    global issues, reservations

    issues.clear()
    reservations.clear()

    # LOAD BORROWED / RETURNED BOOKS
    loans_result = supabase.table("loans").select("*").execute()

    for row in loans_result.data:
        issues.append({
            "id": row.get("id"),
            "book_id": str(row.get("book_id") or ""),
            "book_name": row.get("book_name") or "",
            "student": row.get("student_name") or "",
            "class": row.get("class_name") or "",
            "issue_date": str(row.get("issue_date") or ""),
            "expected_return": str(row.get("expected_return") or ""),
            "return_date": str(row.get("return_date") or ""),
            "returned": row.get("returned", False),
            "status": "returned" if row.get("returned", False) else "borrowed"
        })

    # LOAD RESERVATIONS FROM SUPABASE
    result = supabase.table("reservations").select("*").execute()

    for row in result.data:
        reservation = {
            "id": row.get("id"),
            "book_id": str(row.get("book_id") or ""),
            "book_name": row.get("book_name") or "",
            "student": row.get("student_name") or "",
            "class": row.get("class_name") or "",
            "reserved_on": str(row.get("reserved_on") or ""),
            "reserved_until": str(row.get("reserved_until") or ""),
            "expired": row.get("expired", False)
        }

        reservations.append(reservation)

        # Same reservation displayed in History; no duplicate DB row
        issues.append({
            "reservation_id": reservation["id"],
            "book_id": reservation["book_id"],
            "book_name": reservation["book_name"],
            "student": reservation["student"],
            "class": reservation["class"],
            "issue_date": "---",
            "expected_return": "---",
            "return_date": "---",
            "reserved_on": reservation["reserved_on"],
            "end_date": reservation["reserved_until"],
            "returned": False,
            "status": "expired" if reservation["expired"] else "reserved"
        })





 
books = []
book_counter = 1

classes = []
students = []
issues = []
history = []
reservations = []
UPLOAD_FOLDER = "static/uploads"
timetable_image = None
current_history_view = []

books_from_excel = False
students_from_excel = False

books_excel_name = None
students_excel_name = None

def check_expired_reservations():
    today = date.today()

    try:
        result = (
            supabase.table("reservations")
            .select("*")
            .eq("expired", False)
            .execute()
        )

        for row in result.data:
            expiry = row.get("reserved_until")

            if not expiry:
                continue

            try:
                expiry_date = datetime.strptime(
                    str(expiry)[:10], "%Y-%m-%d"
                ).date()
            except ValueError:
                continue

            if expiry_date < today:
                supabase.table("reservations").update({
                    "expired": True
                }).eq(
                    "id", row["id"]
                ).eq(
                    "expired", False
                ).execute()

        load_history_from_db()

    except Exception as e:
        print("Could not update expired reservations:", e)







# LOGIN
@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        username = request.form.get("username")
        password = request.form.get("password")

        # SIMPLE PASSWORD
        if password == "libmem123":

            session["user"] = username

            return redirect("/dashboard")

    return render_template("login.html")

# LOGOUT
@app.route("/logout")
def logout():

    session.pop("user", None)

    return redirect("/login")


# PROFILE
@app.route("/profile", methods=["GET", "POST"])
def profile():

    if "user" not in session:
        return redirect("/login")

    if request.method == "POST":

        new_name = request.form.get("username")

        if new_name:
            session["user"] = new_name

        return redirect("/dashboard")

    return render_template(
        "profile.html",
        user=session["user"]
    )

    
# =========================
# FORCE DELETE LIBMEM
# =========================

@app.route("/force_delete_libmem", methods=["POST"])
def force_delete_libmem():

    global books, students, classes
    global issues, history, reservations
    global current_history_view, book_counter
    global timetable_image

    if "user" not in session:
        return redirect("/login")

    # BOTH CONFIRMATIONS ARE REQUIRED
    confirmation = request.form.get("confirm_text", "")
    checkbox = request.form.get("confirm_checkbox")

    if confirmation != "DELETE" or checkbox != "yes":
        return """
        <h2>Deletion cancelled</h2>
        <p>The confirmation was incomplete or incorrect.</p>
        <a href="/profile">Return to Profile</a>
        """, 400

    try:
        # DELETE TRANSACTIONS FIRST
        # This helps avoid foreign-key conflicts.
        tables_to_clear = [
            "loans",
            "reservations",
            "students",
            "classes",
            "books"
        ]

        for table_name in tables_to_clear:
            (
                supabase.table(table_name)
                .delete()
                .not_.is_("id", "null")
                .execute()
            )

        # CLEAR THE IN-MEMORY DATA TOO
        books.clear()
        students.clear()
        classes.clear()
        issues.clear()

# Do NOT use history.clear().
# "history" is also the name of your History route.

        reservations.clear()
        current_history_view.clear()



        book_counter = 1
        timetable_image = None

        return """
        <!DOCTYPE html>
        <html>
        <head>
            <title>LiBMeM Reset Complete</title>
            <meta name="viewport" content="width=device-width, initial-scale=1">
            <style>
                body {
                    background: #f8fafc;
                    font-family: Arial, sans-serif;
                    display: grid;
                    place-items: center;
                    min-height: 100vh;
                    margin: 0;
                    color: #172033;
                }
                .card {
                    background: white;
                    padding: 36px;
                    border-radius: 20px;
                    text-align: center;
                    max-width: 420px;
                    margin: 20px;
                    box-shadow: 0 10px 35px #00000012;
                }
                h2 { color: #15803d; }
                a {
                    display: inline-block;
                    margin-top: 16px;
                    padding: 12px 22px;
                    background: #172033;
                    color: white;
                    text-decoration: none;
                    border-radius: 10px;
                }
            </style>
        </head>
        <body>
            <div class="card">
                <h2>LiBMeM has been reset</h2>
                <p>Application records have been cleared.</p>
                <p>You can now begin setting up your library again.</p>
                <a href="/dashboard">Return to Dashboard</a>
            </div>
        </body>
        </html>
        """

    except Exception as e:
        print("FORCE DELETE ERROR:", repr(e))

        return """
        <h2>Reset was not performed </h2>
        <p>Please inform Tanush about this.</p>
        <p>Some tables may have not been cleared.</p>
        <a href="/profile">Return to Profile</a>
        """, 500

    

# DASHBOARD
@app.route("/")
@app.route("/dashboard")
def dashboard():

    check_expired_reservations()

    if "user" not in session:
        return redirect("/login")

    today = date.today()

    # CHECK LATE RETURNS
    late_returns = []

    for i in issues:

        if (

            not i.get("returned")

            and i.get("expected_return")

            and i.get("status") != "reserved"

            and i.get("status") != "expired"

        ):

            try:

                return_date = datetime.strptime(

                    i["expected_return"],
                    "%Y-%m-%d"

                ).date()

                if return_date <= today:

                    late_returns.append(i)

            except:
                pass

    return render_template(

        "dashboard.html",

        late_returns=late_returns,

        user=session["user"],

        timetable_image=timetable_image

    )


   
import os

# TIMETABLE IMAGE
from werkzeug.utils import secure_filename

UPLOAD_FOLDER = "static/uploads"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


@app.route("/upload_timetable", methods=["POST"])
def upload_timetable():

    global timetable_image

    file = request.files.get("image")

    if file and file.filename != "":

        filename = secure_filename(file.filename)

        upload_path = os.path.join(
            UPLOAD_FOLDER,
            filename
        )

        file.save(upload_path)

        timetable_image = "uploads/" + filename

    return redirect("/dashboard")


# DELETE TIMETABLE
@app.route("/delete_timetable")
def delete_timetable():

    global timetable_image

    timetable_image = None

    return redirect("/dashboard")



# ADD BOOKS PAGE
@app.route("/books", methods=["GET", "POST"])
def books_page():

    load_books_from_db()

    if books_from_excel:
        return render_template(
            "manual_blocked.html",
            message="Manual book adding is not allowed while using From Excel."
        )

    global book_counter

    if request.method == "POST":

        name = request.form.get("name")
        author = request.form.get("author")
        publisher = request.form.get("publisher")
        isbn = request.form.get("isbn")
        price = request.form.get("price")
        subtitle = request.form.get("subtitle")
        year_published = request.form.get("year_published")
        cover = request.form.get("cover")


        # ==========================================
        # VALIDATION
        # ==========================================

        if (
            not name
            or not author
            or not isbn
            or not price
            or not year_published
        ):
            return redirect("/books")


        # ==========================================
        # GET NEXT ID FROM SUPABASE
        # ==========================================

        book_id = get_next_book_id()


        # ==========================================
        # CREATE BOOK OBJECT
        # ==========================================

        new_book = {
            "id": book_id,
            "name": name,
            "author": author,
            "publisher": publisher,
            "isbn": isbn,
            "price": price,
            "subtitle": subtitle,
            "year_published": year_published,
            "cover": cover
        }


        # ==========================================
        # SAVE TO SUPABASE FIRST
        # ==========================================

        save_book_to_db(new_book)


        # ==========================================
        # ONLY ADD LOCALLY AFTER DATABASE SUCCESS
        # ==========================================

        books.append(new_book)


        # Keep counter roughly synchronized
        book_counter += 1


        return redirect("/books")


    return render_template(
        "books.html",
        books=books
    )






# EDIT
@app.route("/edit_book/<id>", methods=["GET", "POST"])
def edit_book(id):
    book = next((b for b in books if b["id"] == id), None)

    if request.method == "POST":
        book["name"] = request.form.get("name")
        book["author"] = request.form.get("author")
        book["publisher"] = request.form.get("publisher")
        book["year_published"] = request.form.get("year_published")
        book["isbn"] = request.form.get("isbn")
        book["price"] = request.form.get("price")
        book["subtitle"] = request.form.get("subtitle")
        book["cover"] = request.form.get("cover")

        update_book_in_db(book)

        return redirect("/books")

    return render_template("edit_book.html", book=book)


# SEARCH
@app.route("/search")
def search():
    query = request.args.get("q", "").lower()
    filtered = [b for b in books if b["id"].lower().startswith(query)]
    return render_template("books.html", books=filtered)


# SEARCH BOOKS BY NAME
@app.route("/search_name")
def search_name():
    query = request.args.get("q", "").strip().lower()

    filtered = [
        b for b in books
        if query in b.get("name", "").lower()
    ]

    return render_template("books.html", books=filtered)



#THE STUDENTS PAGE
@app.route("/students", methods=["GET", "POST"])
def students_page():

    load_students_from_db()

    if students_from_excel:
        return render_template(
            "manual_blocked.html",
            message="Manual student and class adding is not allowed while using From Excel."
        )

    if request.method == "POST":

     if request.method == "POST":

      if students_from_excel:

        return render_template(

        "students.html",

        students_from_excel=True

    )



    if request.method == "POST":

        class_name = request.form.get("class").strip()
        student_name = request.form.get("student").strip()


        # EMPTY CHECK
        if not class_name or not student_name:

            return "Please fill all fields"


        # FORMAT CHECK → Example: 3 A
        parts = class_name.split(" ")


        # EXACTLY ONE SPACE
        if len(parts) != 2:

            return "Use format like 3 A (one space only)"


        # FIRST PART MUST BE NUMBER
        if not parts[0].isdigit():

            return "Class must start with a number"


        # SECOND PART MUST BE LETTER
        if not parts[1].isalpha():

            return "Section must be a letter"


        # ONLY ONE LETTER
        if len(parts[1]) != 1:

            return "Section must be one letter only"


        # SAVE CLASS
        if class_name not in classes:

            classes.append(class_name)

                    # SAVE STUDENT
        students.append({
            "name": student_name,
            "class": class_name
        })

        # SAVE TO SUPABASE
        save_class_to_db(class_name)
        save_student_to_db(students[-1])

        return redirect("/students")
    

        
    search = request.args.get("search", "").lower()


    if search:

        filtered_students = [

            s for s in students

            if search in s["name"].lower()
            or search in s["class"].lower()

        ]

    else:

        filtered_students = students


    return render_template(

        "students.html",

        classes=classes,
        students=filtered_students

    )



#DELETE STUDENT
@app.route("/delete_student/<int:index>")
def delete_student(index):
    if 0 <= index < len(students):

        student = students[index]

        delete_student_from_db(
            student["name"],
            student["class"]
        )

        students.pop(index)

    return redirect("/students")




#DELETE CLASS
@app.route("/delete_class/<class_name>")
def delete_class(class_name):

    delete_class_from_db(class_name)

    global students, classes

    classes = [
        c for c in classes
        if c != class_name
    ]

    students = [
        s for s in students
        if s["class"] != class_name
    ]

    return redirect("/students")



# EDIT STUDENT
@app.route("/edit_student/<int:index>", methods=["GET", "POST"])
def edit_student(index):

    if index < 0 or index >= len(students):
        return redirect("/students")

    student = students[index]

    if request.method == "POST":

        old_name = student["name"]
        old_class = student["class"]
        new_name = (request.form.get("name") or "").strip()

        if not new_name:
            return "Please enter a student name."

        try:
            # UPDATE SUPABASE FIRST
            update_student_in_db(
                old_name,
                old_class,
                new_name
            )

            # RELOAD THE LATEST DATABASE DATA
            load_students_from_db()

            return redirect("/students")

        except Exception as e:
            print("Student update failed:", e)
            return "Could not update student in Supabase. Check the Render logs."

    return render_template(
        "edit_student.html",
        student=student,
        index=index
    )


# EDIT CLASS
@app.route("/edit_class/<old_name>", methods=["GET", "POST"])
def edit_class(old_name):

    if request.method == "POST":

        new_name = (request.form.get("name") or "").strip()

        if not new_name:
            return "Please enter class."

        parts = new_name.split(" ")

        if len(parts) != 2:
            return "Use format like 3 A (one space only)."

        if not parts[0].isdigit():
            return "Class must start with a number."

        if not parts[1].isalpha() or len(parts[1]) != 1:
            return "Section must be one letter only."

        if new_name == old_name:
            return redirect("/students")

        if new_name in classes:
            return "That class already exists."

        try:
            # UPDATE STUDENTS AND CLASS IN SUPABASE
            update_class_in_db(old_name, new_name)

            # RELOAD DATABASE DATA
            load_students_from_db()

            return redirect("/students")

        except Exception as e:
            print("Class update failed:", e)
            return "Could not update class."

    return render_template(
        "edit_class.html",
        old_name=old_name
    )



from datetime import datetime
from flask import request, redirect, render_template

#ISSUE
# ISSUE BOOK
@app.route("/issue", methods=["GET", "POST"])
def issue_book():

    if request.method == "POST":

        book_id = request.form.get("book_id")
        class_name = request.form.get("class_name")
        student_name = request.form.get("student_name")
        expected_return = request.form.get("return_date")


        # ==========================================
        # CHECK REQUIRED FIELDS
        # ==========================================

        if not book_id or not class_name or not student_name or not expected_return:
            return "Please fill out all required fields"


        # ==========================================
        # BLOCK PAST DATE
        # ==========================================

        today = date.today()

        selected_date = datetime.strptime(
            expected_return,
            "%Y-%m-%d"
        ).date()

        if selected_date < today:
            return "You cannot issue a book to a past date"


        # ==========================================
        # FIND BOOK
        # ==========================================

        book = next(
            (
                b for b in books
                if str(b["id"]) == str(book_id)
            ),
            None
        )

        if not book:
            return redirect("/issue")


        # ==========================================
        # CHECK RESERVATION
        # ==========================================

        reserved = next(
            (
                r for r in reservations
                if str(r["book_id"]) == str(book_id)
                and not r.get("expired", False)
            ),
            None
        )

        if reserved:
            return (
                f"This book is reserved by "
                f"{reserved['student']} "
                f"from {reserved['class']}"
            )


        # ==========================================
        # CHECK LOCAL ACTIVE ISSUE
        # ==========================================

        already_issued = any(
            str(i["book_id"]) == str(book_id)
            and not i.get("returned", False)
            and i.get("status") not in ["reserved", "expired"]
            for i in issues
        )

        if already_issued:
            return "Book already issued!"


        # ==========================================
        # CHECK SUPABASE ACTIVE ISSUE
        # ==========================================

        existing_db = (
            supabase
            .table("loans")
            .select("id")
            .eq("book_id", book_id)
            .eq("returned", False)
            .execute()
        )

        if existing_db.data:
            return "Book already issued!"


        # ==========================================
        # TODAY
        # ==========================================

        issue_date = datetime.now().strftime(
            "%Y-%m-%d"
        )


        # ==========================================
        # SAVE TO SUPABASE FIRST
        # ==========================================

        supabase.table("loans").insert({
            "book_id": book_id,
            "book_name": book["name"],
            "student_name": student_name,
            "class_name": class_name,
            "issue_date": issue_date,
            "expected_return": expected_return,
            "return_date": None,
            "returned": False
        }).execute()


        # ==========================================
        # SAVE TO LOCAL HISTORY
        # ==========================================

        issues.append({

            "book_id": book_id,
            "book_name": book["name"],
            "class": class_name,
            "student": student_name,
            "issue_date": issue_date,
            "expected_return": expected_return,
            "return_date": "",
            "returned": False

        })


        return redirect("/issue")


    # ==========================================
    # GET REQUEST
    # ==========================================

    normal_issues = [

        i for i in issues

        if i.get("status") != "reserved"
        and i.get("status") != "expired"

    ]

    return render_template(
    "issue.html",
    books=books,
    students=students,
    classes=classes,
    issues=normal_issues,
    has_borrowed=has_borrowed,
    today=date.today().strftime("%Y-%m-%d"),
    seven_days_from_today=(
        date.today() + timedelta(days=7)
    ).strftime("%Y-%m-%d")
)




#BRO




@app.route("/get_students/<class_name>")
def get_students(class_name):
    filtered = [s for s in students if s["class"] == class_name]
    return {"students": filtered}

@app.route("/get_book_info/<book_id>")
def get_book_info(book_id):

    book = next(

        (b for b in books if str(b["id"]) == str(book_id)),

        None

    )

    if not book:

        return {

            "found": False

        }

    return {

        "found": True,

        "name": book["name"],

        "author": book["author"],

        "publisher": book["publisher"],

        "year_published": book["year_published"],

        "isbn": book["isbn"],

        "price": book["price"],

        "subtitle": book["subtitle"]

    }






@app.route("/search_borrowers")
def search_borrowers():

    query = request.args.get("q", "").lower()

    filtered = [

        i for i in issues

        if (

            query in i["student"].lower()

            or query in i["book_id"].lower()

            or query in i["class"].lower()

        )

        and i.get("status") != "reserved"
        and i.get("status") != "expired"

    ]

    return render_template(

        "issue.html",

        books=books,
        students=students,
        classes=classes,

        issues=filtered,

        has_borrowed=has_borrowed

    )


# EDIT EXPECTED RETURN DATE
@app.route("/edit_return_date/<book_id>", methods=["GET", "POST"])
def edit_return_date(book_id):

    issue = next(
        (
            i for i in issues
            if str(i["book_id"]) == str(book_id)
            and not i.get("returned", False)
            and i.get("status") not in ("reserved", "expired")
        ),
        None
    )

    if not issue:
        return redirect("/issue")

    if request.method == "POST":

        new_return_date = request.form.get("return_date")

        if not new_return_date:
            return "Please select an expected return date."

        try:
            datetime.strptime(new_return_date, "%Y-%m-%d")
        except ValueError:
            return "Invalid return date."

        # SAVE THE CHANGE TO SUPABASE FIRST
        supabase.table("loans").update({
            "expected_return": new_return_date
        }).eq(
            "book_id", book_id
        ).eq(
            "returned", False
        ).execute()

        # UPDATE THE IN-MEMORY RECORD
        issue["expected_return"] = new_return_date

        return redirect("/issue")

    return render_template(
        "edit_return_date.html",
        issue=issue
    )


# RESERVATION
@app.route("/reserve", methods=["GET", "POST"])
def reserve_book():

    if request.method == "POST":
        book_id = (request.form.get("book_id") or "").strip()
        class_name = (request.form.get("class_name") or "").strip()
        student_name = (request.form.get("student_name") or "").strip()
        reserved_until = (request.form.get("reserved_until") or "").strip()

        if not all([book_id, class_name, student_name, reserved_until]):
            return "Please fill out all required fields."

        try:
            chosen_date = datetime.strptime(
                reserved_until, "%Y-%m-%d"
            ).date()
        except ValueError:
            return "Please select a valid reservation date."

        if chosen_date < date.today():
            return "You cannot choose a date before today."

        load_books_from_db()
        load_students_from_db()
        check_expired_reservations()

        book = next(
            (b for b in books if str(b["id"]) == book_id),
            None
        )

        if not book:
            return "Book not found."

        valid_student = any(
            s["name"] == student_name and s["class"] == class_name
            for s in students
        )

        if not valid_student:
            return "Please select a valid student from the selected class."

        active_reservation = (
            supabase.table("reservations")
            .select("id, student_name, class_name")
            .eq("book_id", book_id)
            .eq("expired", False)
            .limit(1)
            .execute()
        )

        if active_reservation.data:
            existing = active_reservation.data[0]
            return (
                "This book is already reserved by "
                f"{existing.get('student_name', 'another student')} "
                f"from {existing.get('class_name', 'another class')}."
            )

        active_loan = (
            supabase.table("loans")
            .select("id")
            .eq("book_id", book_id)
            .eq("returned", False)
            .limit(1)
            .execute()
        )

        if active_loan.data:
            return "This book is currently issued and cannot be reserved."

        today = date.today().strftime("%Y-%m-%d")

        # SAVE DIRECTLY TO SUPABASE
        supabase.table("reservations").insert({
            "book_id": book_id,
            "book_name": book["name"],
            "student_name": student_name,
            "class_name": class_name,
            "reserved_on": today,
            "reserved_until": reserved_until,
            "expired": False
        }).execute()

        load_history_from_db()
        return redirect("/reserve")

    load_books_from_db()
    load_students_from_db()
    check_expired_reservations()

    query = (request.args.get("search") or "").strip().lower()

    visible_reservations = [
    r for r in reservations
    if not r.get("expired", False)
    and (
        query in r["book_id"].lower()
        or query in r["student"].lower()
        or query in r["class"].lower()
    )
]



    return render_template(
        "reserve.html",
        classes=classes,
        reservations=visible_reservations,
        today=date.today().strftime("%Y-%m-%d")
    )


@app.route("/search_reservations")
def search_reservations():
    query = (
        request.args.get("q")
        or request.args.get("search")
        or ""
    ).strip().lower()

    load_history_from_db()

    visible_reservations = [
    r for r in reservations
    if not r.get("expired", False)
    and (
        query in r["book_id"].lower()
        or query in r["student"].lower()
        or query in r["class"].lower()
    )
]
    

    return render_template(
        "reserve.html",
        classes=classes,
        reservations=visible_reservations,
        today=date.today().strftime("%Y-%m-%d")
    )


@app.route("/edit_reservation/<book_id>", methods=["GET", "POST"])
def edit_reservation(book_id):
    load_history_from_db()

    reservation = next(
        (
            r for r in reversed(reservations)
            if str(r["book_id"]) == str(book_id)
            and not r.get("expired", False)
        ),
        None
    )

    if not reservation or not reservation.get("id"):
        return redirect("/reserve")

    if request.method == "POST":
        new_date = (
            request.form.get("reserved_until") or ""
        ).strip()

        try:
            parsed_date = datetime.strptime(
                new_date, "%Y-%m-%d"
            ).date()
        except ValueError:
            return "Please select a valid expiry date."



        
        supabase.table("reservations").update({
            "reserved_until": new_date,
            "expired": False
        }).eq("id", reservation["id"]).execute()

        load_history_from_db()
        return redirect("/reserve")

    return render_template(
        "edit_reservation.html",
        reservation=reservation
    )




    






#RETURN MAIN PAGE
@app.route("/return")
def return_page():

    normal_issues = [

        i for i in issues

        if not i.get("returned", False)

        and i.get("status") != "reserved"

        and i.get("status") != "expired"

    ]

    return render_template(

        "return.html",

        issues=normal_issues,

        classes=classes,

        books=books

    )




# RETURN BY BOOK ID

@app.route("/return_book", methods=["POST"])
def return_book_post():
    book_id = request.form.get("book_id")

    if not book_id:
        return redirect("/return")

    for i in issues:

        if (
            i["book_id"] == book_id
            and not i["returned"]
            and i.get("status") != "reserved"
            and i.get("status") != "expired"
        ):

            # MARK AS RETURNED
            i["returned"] = True

            # SAVE RETURN DATE
            i["return_date"] = datetime.now().strftime("%Y-%m-%d")

            # UPDATE SUPABASE
            supabase.table("loans").update({
                "returned": True,
                "return_date": i["return_date"]
            }).eq("book_id", book_id).eq("returned", False).execute()

            break

    return redirect("/return")






# RETURN BY CLASS (MULTIPLE STUDents)
@app.route("/return_class", methods=["POST"])
def return_class():

    class_name = request.form.get("class_name")
    selected_students = request.form.getlist("students")

    if not class_name or not selected_students:
        return redirect("/return")

    for i in issues:

      if (
        i["class"] == class_name
        and i["student"] in selected_students
        and not i["returned"]
        and i.get("status") != "reserved"
        and i.get("status") != "expired"
    ):

        # MARK RETURNED
        i["returned"] = True

        # SAVE RETURN DATE
        i["return_date"] = datetime.now().strftime("%Y-%m-%d")

        # UPDATE SUPABASE
        supabase.table("loans").update({
            "returned": True,
            "return_date": i["return_date"]
        }).eq("book_id", i["book_id"]).eq("returned", False).execute()


    
    
    return redirect("/return")


# HISTORY
@app.route("/history")
def history():

    global current_history_view

    # Make sure reservation expiry is checked first
    check_expired_reservations()

    # ==========================================
    # BUILD HISTORY VIEW
    # ==========================================

    history_rows = []

    # ------------------------------------------
    # BORROWED / RETURNED LOANS
    # ------------------------------------------

    for i in issues:

        # Do not show expired/reserved entries here
        if i.get("status") in ["reserved", "expired"]:
            continue

        history_rows.append(i)


    # ------------------------------------------
    # RESERVED / EXPIRED RESERVATIONS
    # ------------------------------------------

    for i in issues:

        if i.get("status") == "reserved":
            history_rows.append(i)

        elif i.get("status") == "expired":
            history_rows.append(i)


    # Save exactly what is currently displayed
    current_history_view = history_rows


    # ==========================================
    # SEARCH VALUES
    # ==========================================

    search_student = request.args.get(
        "student", ""
    ).lower()

    search_class = request.args.get(
        "class", ""
    ).lower()

    search_book = request.args.get(
        "book", ""
    ).lower()

    search_borrow = request.args.get(
        "borrow_date", ""
    )

    search_return = request.args.get(
        "return_date", ""
    )

    search_actual = request.args.get(
        "actual_return", ""
    )

    search_reserve = request.args.get(
        "reserve_date", ""
    )

    search_expiry = request.args.get(
        "expiry_date", ""
    )

    search_status = request.args.get(
        "status", ""
    ).lower()

    search_reservation = request.args.get(
        "reservation_status", ""
    ).lower()


    # ==========================================
    # FILTER
    # ==========================================

    filtered = history_rows


    # STUDENT
    if search_student:

        filtered = [
            i for i in filtered
            if search_student in
            i.get("student", "").lower()
        ]


    # CLASS
    if search_class:

        filtered = [
            i for i in filtered
            if search_class in
            i.get("class", "").lower()
        ]


    # BOOK
    if search_book:

        filtered = [
            i for i in filtered
            if search_book in
            str(i.get("book_id", "")).lower()
        ]


    # BORROW DATE
    if search_borrow:

        filtered = [
            i for i in filtered
            if search_borrow in
            str(i.get("issue_date", ""))
        ]


    # RESERVE DATE
    if search_reserve:

        filtered = [
            i for i in filtered
            if search_reserve in
            str(i.get("reserved_on", ""))
        ]


    # EXPECTED RETURN
    if search_return:

        filtered = [
            i for i in filtered
            if search_return in
            str(i.get("expected_return", ""))
        ]


    # ACTUAL RETURN
    if search_actual:

        filtered = [
            i for i in filtered
            if search_actual in
            str(i.get("return_date", ""))
        ]


    # EXPIRY DATE
    if search_expiry:

        filtered = [
            i for i in filtered
            if search_expiry in
            str(i.get("end_date", ""))
        ]


    # ==========================================
    # BORROWED / RETURNED
    # ==========================================

    if search_status:

        if search_status == "returned":

            filtered = [
                i for i in filtered
                if i.get("returned", False)
                and i.get("status") != "reserved"
                and i.get("status") != "expired"
            ]


        elif search_status == "borrowed":

            filtered = [
                i for i in filtered
                if not i.get("returned", False)
                and i.get("status") != "reserved"
                and i.get("status") != "expired"
            ]


    # ==========================================
    # RESERVED / EXPIRED
    # ==========================================

    if search_reservation:

        if search_reservation == "reserved":

            filtered = [
                i for i in filtered
                if i.get("status") == "reserved"
            ]


        elif search_reservation == "expired":

            filtered = [
                i for i in filtered
                if i.get("status") == "expired"
            ]


    # ==========================================
    # RENDER HISTORY
    # ==========================================

    return render_template(
        "history.html",
        issues=filtered,
        books=books,
        classes=classes
    )


    from datetime import datetime, timedelta

#DELETE HISTORY
# DELETE HISTORY
@app.route("/delete_history", methods=["POST"])
def delete_history():

    option = request.form.get("delete_option")

    if not option:
        return redirect("/history")

    global issues

    # ==========================================
    # DELETE FROM SUPABASE FIRST
    # ==========================================

    delete_history_from_db(option)


    # ==========================================
    # DELETE FROM FLASK MEMORY
    # ==========================================

    from datetime import datetime, timedelta


    # ------------------------------------------
    # DELETE ALL COMPLETED HISTORY
    # ------------------------------------------

    if option == "all":

        issues = [
            i for i in issues

            # KEEP ACTIVE BORROWED
            if (
                not i.get("returned", False)
                and i.get("status") != "reserved"
                and i.get("status") != "expired"
            )

            # KEEP ACTIVE RESERVED
            or i.get("status") == "reserved"
        ]

        return redirect("/history")


    # ------------------------------------------
    # TIME-BASED DELETE
    # ------------------------------------------

    days_map = {
        "5": 5,
        "10": 10,
        "20": 20,
        "30": 30
    }

    days = days_map.get(option)

    if not days:
        return redirect("/history")


    cutoff = datetime.now() - timedelta(days=days)

    new_issues = []


    for i in issues:

        # ======================================
        # ACTIVE BORROWED → KEEP
        # ======================================

        if (
            not i.get("returned", False)
            and i.get("status") != "reserved"
            and i.get("status") != "expired"
        ):

            new_issues.append(i)
            continue


        # ======================================
        # ACTIVE RESERVED → KEEP
        # ======================================

        if i.get("status") == "reserved":

            new_issues.append(i)
            continue


        # ======================================
        # RETURNED
        # ======================================

        if i.get("returned", False):

            try:

                return_date = datetime.strptime(
                    i["return_date"],
                    "%Y-%m-%d"
                )

                if return_date > cutoff:
                    new_issues.append(i)

            except:

                new_issues.append(i)


        # ======================================
        # EXPIRED RESERVATION
        # ======================================

        elif i.get("status") == "expired":

            try:

                expiry_date = datetime.strptime(
                    i["end_date"],
                    "%Y-%m-%d"
                )

                if expiry_date > cutoff:
                    new_issues.append(i)

            except:

                new_issues.append(i)


    issues = new_issues

    return redirect("/history")



@app.route("/history_range_class")
def history_range_class():
    global current_history_view

    from_class = request.args.get("from_class", "").strip()
    to_class = request.args.get("to_class", "").strip()

    if not from_class or not to_class:
        return redirect("/history")

    sorted_classes = sorted(set(str(c) for c in classes))

    if from_class not in sorted_classes or to_class not in sorted_classes:
        return redirect("/history")

    start = sorted_classes.index(from_class)
    end = sorted_classes.index(to_class)

    # Allow the user to select either direction
    if start > end:
        start, end = end, start

    allowed_classes = sorted_classes[start:end + 1]

    filtered = [
        i for i in issues
        if str(i.get("class", "")) in allowed_classes
    ]

    current_history_view = filtered

    return render_template(
        "history.html",
        issues=filtered,
        books=books,
        classes=classes
    )


@app.route("/history_range_date")
def history_range_date():
    global current_history_view

    from_date = request.args.get("from_date", "").strip()
    to_date = request.args.get("to_date", "").strip()

    if not from_date or not to_date:
        return redirect("/history")

    # Validate both dates
    try:
        start_date = datetime.strptime(from_date, "%Y-%m-%d").date()
        end_date = datetime.strptime(to_date, "%Y-%m-%d").date()
    except ValueError:
        return redirect("/history")

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    filtered = []

    for i in issues:
        # Reservations use their reservation date.
        # Borrowed/returned books use their issue date.
        if i.get("status") in ("reserved", "expired"):
            date_to_check = i.get("reserved_on")
        else:
            date_to_check = i.get("issue_date")

        if not date_to_check or date_to_check == "---":
            continue

        try:
            record_date = datetime.strptime(
                str(date_to_check)[:10], "%Y-%m-%d"
            ).date()
        except (ValueError, TypeError):
            continue

        if start_date <= record_date <= end_date:
            filtered.append(i)

    current_history_view = filtered

    return render_template(
        "history.html",
        issues=filtered,
        books=books,
        classes=classes
    )

   
@app.route("/export_history")
def export_history():

    global current_history_view

    wb = Workbook()
    ws = wb.active

    ws.title = "History"

    ws.append([

        "Student",
        "Class",
        "Book ID",
        "Borrow Date",
        "Reserve Date",
        "Expected Return",
        "Actual Return",
        "Expiry Date",
        "Status"

    ])


    for i in current_history_view:

        if i.get("status") == "reserved":

            status = "RESERVED"

        elif i.get("status") == "expired":

            status = "EXPIRED"

        elif i.get("returned"):

            status = "RETURNED"

        else:

            status = "BORROWED"


        if (

            i.get("status") == "reserved"

            or

            i.get("status") == "expired"

        ):

            borrow_date = "---"

            reserve_date = i.get(

                "reserved_on",

                "---"

            )

            expected_return = "---"

            actual_return = "---"

            expiry_date = i.get(

                "end_date",

                "---"

            )

        else:

            borrow_date = i.get(

                "issue_date",

                "---"

            )

            reserve_date = "---"

            expected_return = i.get(

                "expected_return",

                "---"

            )

            actual_return = i.get(

                "return_date",

                "Not Returned"

            )

            expiry_date = "---"


        ws.append([

            i.get("student", ""),

            i.get("class", ""),

            i.get("book_id", ""),

            borrow_date,

            reserve_date,

            expected_return,

            actual_return,

            expiry_date,

            status

        ])


    # ONLY AFTER LOOP ENDS

    excel_file = BytesIO()

    wb.save(excel_file)

    excel_file.seek(0)

    return send_file(

        excel_file,

        as_attachment=True,

        download_name="history_export.xlsx",

        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )


# =========================
# EXCEL IMPORT PAGE
# =========================

@app.route("/excel_import")
def excel_import():

    # Refresh everything from Supabase before displaying the page
    load_books_from_db()
    load_students_from_db()

    return render_template(
        "excel_import.html",
        books_from_excel=books_from_excel,
        students_from_excel=students_from_excel,
        books_excel_name=books_excel_name,
        students_excel_name=students_excel_name,
        books=books,
        students=students,
        classes=classes
    )







# =========================
# UPLOAD BOOKS FROM EXCEL
# =========================

@app.route("/upload_books_excel", methods=["POST"])
def upload_books_excel():
    
    global books
    global books_from_excel
    global books_excel_name
    global book_counter

   
    file = request.files.get("books_file")

    if not file or file.filename == "":
        return """
        <h2 style="color:red;">
        Oops! Looks like you made a mistake.
        </h2>

        <p>No Excel file was selected.</p>

        <a href="/excel_import">Go back</a>
        """

    try:

        wb = load_workbook(file)

        ws = wb.active

        # =========================
        # CHECK EXACT HEADERS
        # =========================

        required_headers = [
            "Book ID",
            "Book Name",
            "Author",
            "Publisher",
            "Year Published",
            "ISBN",
            "Price",
            "Subtitle",
            "Cover"
        ]

        actual_headers = [
            cell.value
            for cell in ws[1]
        ]

        if actual_headers != required_headers:

            return """
            <h2 style="color:red;">
            Oops! Looks like you made a mistake in the Excel format.
            </h2>

            <h3>Required Books Excel format:</h3>

            <p>Column A → Book ID (S0001)</p>
            <p>Column B → Book Name</p>
            <p>Column C → Author</p>
            <p>Column D → Publisher</p>
            <p>Column E → Year Published</p>
            <p>Column F → ISBN</p>
            <p>Column G → Price</p>
            <p>Column H → Subtitle</p>
            <p>Column I → Cover</p>

            <br>

            <p>
            The column headings must be in the exact order shown above.
            </p>

            <a href="/excel_import">Go back</a>
            """

        # =========================
        # CHECK FOR EXTRA COLUMNS
        # =========================

        for row in ws.iter_rows(min_row=2, values_only=True):

            if len(row) > 9:

                return """
                <h2 style="color:red;">
                Oops! Looks like you made a mistake in the Excel format.
                </h2>

                <p>
                Your file contains more than 9 columns.
                </p>

                <p>
                The Books Excel file must contain exactly 9 columns.
                </p>

                <a href="/excel_import">Go back</a>
                """

        # =========================
        # READ BOOKS FIRST
        # =========================

        imported_books = []

        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            # Skip completely empty rows

            if not any(row):

                continue

            book_id = row[0]
            name = row[1]
            author = row[2]
            publisher = row[3]
            year_published = row[4]
            isbn = row[5]
            price = row[6]
            subtitle = row[7]
            cover = row[8]

            # Required fields

            if not book_id or not name or not author:

                return """
                <h2 style="color:red;">
                Oops! Looks like you made a mistake in the Excel format.
                </h2>

                <p>
                Book ID, Book Name and Author cannot be empty.
                </p>

                <a href="/excel_import">Go back</a>
                """

            imported_books.append({

                "id": str(book_id).strip(),

                "name": name,

                "author": author,

                "publisher": publisher,

                "year_published": year_published,

                "isbn": isbn,

                "price": price,

                "subtitle": subtitle if subtitle else "",

                "cover": cover if cover else ""

            })

        # =========================
        # REPLACE OLD BOOK DATA
        # =========================

                # SAVE THE IMPORT TO SUPABASE
        db_books = [
            {
                **book,
                "source": "excel"
            }
            for book in imported_books
        ]

        supabase.table("books").delete().neq(
            "id", "__LIBMEM_NO_MATCHING_BOOK_ID__"
        ).execute()

        supabase.table("books").insert(db_books).execute()

        books = imported_books
        book_counter = 1
        books_from_excel = True
        books_excel_name = secure_filename(file.filename) or file.filename

        return redirect("/excel_import")

    






    except Exception:

        return """
        <h2 style="color:red;">
        Oops! Looks like you uploaded the wrong file.
        </h2>

        <p>
        Please make sure you selected a valid Excel (.xlsx) file
        containing the required Books format.
        </p>

        <h3>Required format:</h3>

        <p>Column A → Book ID (S0001)</p>
        <p>Column B → Book Name</p>
        <p>Column C → Author</p>
        <p>Column D → Publisher</p>
        <p>Column E → Year Published</p>
        <p>Column F → ISBN</p>
        <p>Column G → Price</p>
        <p>Column H → Subtitle</p>
        <p>Column I → Cover</p>

        <a href="/excel_import">Go back</a>
        """


# =========================
# UPLOAD STUDENTS FROM EXCEL
# =========================

@app.route("/upload_students_excel", methods=["POST"])
def upload_students_excel():

    global students
    global classes

    global students_from_excel
    global students_excel_name

    file = request.files.get("students_file")

    if not file or file.filename == "":

        return """
        <h2 style="color:red;">
        Oops! Looks like you made a mistake.
        </h2>

        <p>No Excel file was selected.</p>

        <a href="/excel_import">Go back</a>
        """

    try:

        wb = load_workbook(file)

        imported_students = []

        imported_classes = []

        # =========================
        # CHECK EVERY SHEET
        # =========================

        for sheet_name in wb.sheetnames:

            ws = wb[sheet_name]

            class_name = sheet_name.strip()

            if not class_name:

                return """
                <h2 style="color:red;">
                Oops! Looks like you made a mistake in the Excel format.
                </h2>

                <p>
                Every sheet must have a class name.
                </p>

                <a href="/excel_import">Go back</a>
                """

            imported_classes.append(class_name)

            # =========================
            # CHECK EVERY ROW
            # =========================

            for row in ws.iter_rows(
                values_only=True
            ):

                # Completely empty row

                if not any(row):

                    continue

                # MORE THAN ONE COLUMN

                if len(row) != 1:

                    return """
                    <h2 style="color:red;">
                    Oops! Looks like you made a mistake in the Excel format.
                    </h2>

                    <p>
                    Student sheets must contain student names only.
                    </p>

                    <p>
                    No column headers and no extra columns are allowed.
                    </p>

                    <h3>Example:</h3>

                    <p>Sheet name → 3 A</p>

                    <p>Ali</p>
                    <p>Neharika</p>
                    <p>John</p>

                    <a href="/excel_import">Go back</a>
                    """

                student_name = row[0]

                if not isinstance(student_name, str):

                    return """
                    <h2 style="color:red;">
                    Oops! Looks like you made a mistake in the Excel format.
                    </h2>

                    <p>
                    Student names must be text.
                    </p>

                    <a href="/excel_import">Go back</a>
                    """

                student_name = student_name.strip()

                if not student_name:

                    continue

                imported_students.append({

                    "name": student_name,

                    "class": class_name

                })

        # =========================
        # REPLACE OLD DATA
        # =========================

                # REPLACE OLD DATA IN SUPABASE

        db_students = [
            {
                "name": student["name"],
                "class_name": student["class"],
                "source": "excel"
            }
            for student in imported_students
        ]

        # Delete old students first
        supabase.table("students").delete().neq(
            "name", "__LIBMEM_NO_MATCHING_STUDENT__"
        ).execute()

        # Delete old classes
        supabase.table("classes").delete().neq(
            "name", "__LIBMEM_NO_MATCHING_CLASS__"
        ).execute()

        # Insert imported classes
        supabase.table("classes").insert([
            {
                "name": class_name,
                "source": "excel"
            }
            for class_name in imported_classes
        ]).execute()

        # Insert imported students
        if db_students:
            supabase.table("students").insert(
                db_students
            ).execute()

        # Update in-memory data
        students = imported_students
        classes = imported_classes

        students_from_excel = True
        students_excel_name = file.filename

        return redirect("/excel_import")







    except Exception:

        return """
        <h2 style="color:red;">
        Oops! Looks like you uploaded the wrong file.
        </h2>

        <p>
        Please make sure you selected the Students + Classes Excel file.
        </p>

        <h3>Required format:</h3>

        <p>
        Each sheet name = Class Name
        </p>

        <p>
        Example:
        </p>

        <p>3 A</p>
        <p>3 B</p>

        <br>

        <p>
        Inside each sheet:
        </p>

        <p>
        Student names only
        </p>

        <p>
        No column headers
        </p>

        <p>
        Only one column is allowed.
        </p>

        <a href="/excel_import">Go back</a>
        """


# =========================
# DELETE BOOK EXCEL
# =========================

@app.route("/delete_books_excel")
def delete_books_excel():

    global books
    global books_from_excel
    global books_excel_name
    global book_counter

    supabase.table("books").delete().neq(
    "id", "__LIBMEM_NO_MATCHING_BOOK_ID__"
).execute()

    

    books = []

    book_counter = 1

    books_from_excel = False

    books_excel_name = None

    return redirect("/excel_import")


# =========================
# DELETE STUDENT EXCEL
# =========================

@app.route("/delete_students_excel")
def delete_students_excel():

    global students
    global classes

    global students_from_excel
    global students_excel_name

    supabase.table("students").delete().neq(
    "name", "__LIBMEM_NO_MATCHING_STUDENT__"
).execute()

    supabase.table("classes").delete().neq(
    "name", "__LIBMEM_NO_MATCHING_CLASS__"
).execute()



    students = []

    classes = []

    students_from_excel = False

    students_excel_name = None

    return redirect("/excel_import")

@app.errorhandler(404)
def page_not_found(error):

    return render_template(
        "error.html",
        error_code=404
    ), 404


@app.errorhandler(405)
def method_not_allowed(error):

    return render_template(
        "error.html",
        error_code=405
    ), 405


@app.errorhandler(403)
def forbidden(error):

    return render_template(
        "error.html",
        error_code=403
    ), 403


@app.errorhandler(500)
def internal_server_error(error):

    return render_template(
        "error.html",
        error_code=500
    ), 500


# LOAD DATABASE DATA
try:
    load_books_from_db()
except Exception as e:
    print("Could not load books from Supabase:", e)

    # =========================
# STUDENT / CLASS HELPERS
# =========================

def load_students_from_db():
    global students, classes
    global students_from_excel, students_excel_name

    student_result = supabase.table("students").select("*").execute()
    class_result = supabase.table("classes").select("*").execute()

    student_rows = student_result.data or []
    class_rows = class_result.data or []

    students.clear()
    classes.clear()

    # Restore Excel mode from Supabase
    all_rows = student_rows + class_rows

    students_from_excel = (
        bool(all_rows)
        and all(row.get("source") == "excel" for row in all_rows)
    )

    students_excel_name = (
        "Excel data saved in Supabase"
        if students_from_excel else None
    )

    for row in student_rows:
        students.append({
            "name": row["name"],
            "class": row["class_name"]
        })

    for row in class_rows:
        classes.append(row["name"])




def save_class_to_db(class_name):
    existing = (
        supabase.table("classes")
        .select("id")
        .eq("name", class_name)
        .execute()
    )

    if not existing.data:
        supabase.table("classes").insert({
            "name": class_name,
            "source": "manual"
        }).execute()


def save_student_to_db(student):
    supabase.table("students").insert({
        "name": student["name"],
        "class_name": student["class"],
        "source": "manual"
    }).execute()


def delete_student_from_db(name, class_name):
    supabase.table("students").delete().eq(
        "name", name
    ).eq(
        "class_name", class_name
    ).execute()


def delete_class_from_db(class_name):
    supabase.table("students").delete().eq(
        "class_name", class_name
    ).execute()

    supabase.table("classes").delete().eq(
        "name", class_name
    ).execute()


def update_student_in_db(old_name, class_name, new_name):
    supabase.table("students").update({
        "name": new_name
    }).eq(
        "name", old_name
    ).eq(
        "class_name", class_name
    ).execute()


try:
    load_students_from_db()
except Exception as e:
    print("Could not load students from Supabase:", e)

try:
    load_history_from_db()
except Exception as e:
    print("Could not load history from Supabase:", e)






if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)